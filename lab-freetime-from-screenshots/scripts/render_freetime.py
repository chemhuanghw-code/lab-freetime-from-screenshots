#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""课表截图 -> 人员空闲表 渲染器

把「人 × 星期 × 大节」的有课数据渲染成实验室人员空闲表（HTML / XLSX / CSV / PNG）。

用法:
  python render_freetime.py data.json
  python render_freetime.py data.json -o out --week 8.31-9.6 --monday 8.31
  python render_freetime.py --template
"""
import argparse
import csv
import datetime
import json
import os
import sys

DEFAULT_SLOTS = ["上午 1.2 节", "上午 3.4 节", "下午 5.6 节", "下午 7.8 节", "晚上"]
DEFAULT_DAYS = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]
DAY_SHORT = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]

GREEN_HEAD = "D9EFDD"
GREEN_HEAD2 = "EAF6EC"
GREEN_TEXT = "1B6B2C"
WARN_BG = "FFF4E5"
WARN_TEXT = "B06F00"
LINE = "B9C2C9"


# ---------------------------------------------------------------- 数据准备
def load_data(path, args):
    with open(path, "r", encoding="utf-8-sig") as f:
        raw = json.load(f)

    meta = raw.get("meta", {}) or {}
    slots = raw.get("slots") or DEFAULT_SLOTS
    days = raw.get("days") or DEFAULT_DAYS
    people = raw.get("people") or []

    if args.week:
        meta["week_label"] = args.week
    if args.monday:
        meta["monday_date"] = args.monday
    meta.setdefault("title", "实验室人员空闲表")
    meta.setdefault("week_label", "")
    meta.setdefault("monday_date", "")

    norm = []
    for p in people:
        name = (p.get("name") or "").strip()
        if not name:
            continue
        busy = set()
        for item in p.get("busy") or []:
            try:
                d, s = int(item[0]), int(item[1])
            except (TypeError, ValueError, IndexError):
                continue
            if 1 <= d <= len(days) and 1 <= s <= len(slots):
                busy.add((d, s))
        norm.append({"name": name, "busy": busy})

    return meta, slots, days, norm


def week_dates(monday_str, days):
    """把 '8.31' 解析成周一日期，返回 7 个 'M.D' 字符串。"""
    if not monday_str:
        return [""] * len(days)
    s = str(monday_str).strip().replace("/", ".").replace("-", ".").replace("月", ".").replace("日", "")
    parts = [x for x in s.split(".") if x != ""]
    if len(parts) < 2:
        return [""] * len(days)
    try:
        mon = datetime.date(datetime.date.today().year, int(parts[0]), int(parts[1]))
    except ValueError:
        return [""] * len(days)
    return [f"{(mon + datetime.timedelta(days=i)).month}.{(mon + datetime.timedelta(days=i)).day}"
            for i in range(len(days))]


def free_map(people, days, slots):
    """返回 grid[slot_idx][day_idx] = [空闲人名, ...]"""
    grid = []
    for si in range(len(slots)):
        row = []
        for di in range(len(days)):
            row.append([p["name"] for p in people if (di + 1, si + 1) not in p["busy"]])
        grid.append(row)
    return grid


def check_anomalies(people, days, slots, grid):
    out = []
    total = len(people)
    if total >= 3:
        for si, row in enumerate(grid):
            for di, names in enumerate(row):
                if len(names) == 0:
                    out.append(f"{days[di]} {slots[si]}：全员有课（0 人空闲）")
    if total >= 5:
        for si, row in enumerate(grid):
            for di, names in enumerate(row):
                if 0 < len(names) <= max(1, total // 5):
                    out.append(f"{days[di]} {slots[si]}：仅 {len(names)} 人空闲（偏少）")
    for p in people:
        n = len(p["busy"])
        if n == 0:
            out.append(f"{p['name']}：整周无任何有课时段 —— 疑似漏读，请核对截图")
        elif n == len(days) * len(slots) and len(days) * len(slots) > 0:
            out.append(f"{p['name']}：整周每个时段都有课 —— 疑似误读，请核对截图")
    return out


# ---------------------------------------------------------------- HTML
def render_html(meta, slots, days, people, grid, dates, path):
    total = len(people)
    h = []
    h.append("<!DOCTYPE html>")
    h.append('<html lang="zh-CN"><head><meta charset="utf-8">')
    h.append(f"<title>{meta['title']}</title>")
    h.append("<style>")
    h.append('body{margin:0;padding:24px;background:#fff;color:#1f2328;'
             'font:14px/1.6 -apple-system,BlinkMacSystemFont,"Microsoft YaHei",sans-serif}')
    h.append("h1{font-size:18px;margin:0 0 6px}")
    h.append(".meta{color:#6b7280;font-size:13px;margin-bottom:14px}")
    h.append("table{border-collapse:collapse;width:100%}")
    h.append(f"th,td{{border:1px solid #{LINE};padding:6px 8px;text-align:center;"
             f"vertical-align:middle;font-size:13px}}")
    h.append(f"thead th{{background:#{GREEN_HEAD};text-align:center;font-weight:600}}")
    h.append(f"thead tr:nth-child(2) th{{background:#{GREEN_HEAD2};font-weight:500;"
             f"color:#48505a;font-size:12px}}")
    h.append(f"td.slot{{background:#{GREEN_HEAD};font-weight:600;text-align:center;"
             f"white-space:nowrap;width:96px}}")
    h.append(".n{line-height:1.8;min-height:20px}")
    h.append(f".c{{margin-top:3px;font-weight:700;color:#{GREEN_TEXT};font-size:13px}}")
    h.append(f"td.zero{{background:#{WARN_BG}}}")
    h.append(f"td.zero .c{{color:#{WARN_TEXT}}}")
    h.append(".legend{margin-top:12px;color:#6b7280;font-size:12px}")
    h.append("@media print{@page{size:A4 landscape;margin:10mm}body{padding:0}}")
    h.append("</style></head><body>")
    h.append(f"<h1>{meta['title']}</h1>")
    h.append(f'<div class="meta">周期 {meta["week_label"]} ｜ 共 {total} 人 ｜ '
             f'粒度：{len(slots)} 个大节</div>')
    h.append("<table><thead>")
    h.append("<tr><th>周期</th>" + "".join(f"<th>{d}</th>" for d in days) + "</tr>")
    h.append("<tr><th>日期</th>" + "".join(f"<th>{d}</th>" for d in dates) + "</tr>")
    h.append("</thead><tbody>")
    for si, slot in enumerate(slots):
        h.append(f'<tr><td class="slot">{slot}</td>')
        for di in range(len(days)):
            names = grid[si][di]
            cls = ' class="zero"' if not names else ""
            body = "、".join(names) if names else "（全员有课）"
            h.append(f'<td{cls}><div class="n">{body}</div>'
                     f'<div class="c">{len(names)}</div></td>')
        h.append("</tr>")
    h.append("</tbody></table>")
    h.append('<div class="legend">数字为该时段空闲人数；黄色底 = 该时段无人空闲。</div>')
    h.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(h))


# ---------------------------------------------------------------- CSV
def render_csv(meta, slots, days, people, grid, dates, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["周期", meta["week_label"]])
        w.writerow(["日期"] + dates)
        w.writerow(["星期"] + days)
        for si, slot in enumerate(slots):
            row = [slot]
            for di in range(len(days)):
                names = grid[si][di]
                row.append(("、".join(names) + f"（{len(names)}）") if names else "（全员有课）0")
            w.writerow(row)


# ---------------------------------------------------------------- XLSX
def render_xlsx(meta, slots, days, people, grid, dates, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color=LINE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=GREEN_HEAD)
    head2_fill = PatternFill("solid", fgColor=GREEN_HEAD2)
    warn_fill = PatternFill("solid", fgColor=WARN_BG)
    # 所有单元格一律水平+垂直居中
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "人员空闲表"
    ncol = 1 + len(days)

    for i, wd in enumerate([16] + [20] * len(days), start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wd

    for c in range(1, ncol + 1):          # 合并前先给整行打上居中，避免合并区右侧留"常规"格式
        ws.cell(row=1, column=c).alignment = center
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    t = ws.cell(row=1, column=1, value=f"{meta['title']}    周期 {meta['week_label']}    "
                                       f"共 {len(people)} 人")
    t.font = Font(size=13, bold=True)
    t.alignment = center
    ws.row_dimensions[1].height = 26

    ws.cell(row=2, column=1, value="周期")
    for j, d in enumerate(days):
        ws.cell(row=2, column=2 + j, value=d)
    ws.cell(row=3, column=1, value="日期")
    for j, d in enumerate(dates):
        ws.cell(row=3, column=2 + j, value=d)

    for r in (2, 3):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center
            cell.fill = head_fill if r == 2 else head2_fill
            cell.font = Font(bold=(r == 2), size=11 if r == 2 else 10)
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 20

    r0 = 4
    for si, slot in enumerate(slots):
        r = r0 + si
        sc = ws.cell(row=r, column=1, value=slot)
        sc.border = border
        sc.alignment = center
        sc.fill = head_fill
        sc.font = Font(bold=True)
        max_lines = 1
        for di in range(len(days)):
            names = grid[si][di]
            txt = "、".join(names) if names else "（全员有课）"
            cell = ws.cell(row=r, column=2 + di, value=f"{txt}\n{len(names)}")
            cell.border = border
            cell.alignment = center
            if not names:
                cell.fill = warn_fill
                cell.font = Font(color=WARN_TEXT, bold=True)
            lines = max(1, (len(txt) + 13) // 14) + 1
            max_lines = max(max_lines, lines)
        ws.row_dimensions[r].height = max(34, max_lines * 16)
    wb.save(path)


# ---------------------------------------------------------------- PNG
def wrap_names(names, limit=12):
    lines, cur = [], ""
    for n in names:
        add = n if cur == "" else "、" + n
        if cur and len(cur) + len(add) > limit:
            lines.append(cur)
            cur = n
        else:
            cur += add
    if cur:
        lines.append(cur)
    return lines or [""]


def render_png(meta, slots, days, people, grid, dates, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import Rectangle

    avail = {f.name for f in font_manager.fontManager.ttflist}
    for cand in ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC", "SimSun"]:
        if cand in avail:
            plt.rcParams["font.sans-serif"] = [cand]
            break
    plt.rcParams["axes.unicode_minus"] = False

    label_w, day_w = 1.35, 2.05
    total_w = label_w + day_w * len(days)
    pad, line_h = 0.09, 0.205
    hdr_h = 0.34
    title_h = 0.52

    row_h = []
    counts = []
    for si in range(len(slots)):
        mx = 1
        for di in range(len(days)):
            mx = max(mx, len(wrap_names(grid[si][di])))
        row_h.append(max(0.62, mx * line_h + 0.34))
        counts.append([len(grid[si][di]) for di in range(len(days))])

    total_h = title_h + hdr_h * 2 + sum(row_h) + 0.18
    fig = plt.figure(figsize=(total_w, total_h), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, total_w)
    ax.set_ylim(0, total_h)
    ax.invert_yaxis()
    ax.axis("off")

    ax.text(pad, title_h * 0.5, f"{meta['title']}　周期 {meta['week_label']}　共 {len(people)} 人",
            fontsize=12, va="center", ha="left", weight="bold")

    xs = [0.0]
    xs.append(label_w)
    for _ in days:
        xs.append(xs[-1] + day_w)

    def cell(x, y, w, h, txt, fc, fs=9.5, bold=False, color="#1f2328", ha="left", va="top"):
        ax.add_patch(Rectangle((x, y), w, h, facecolor=fc, edgecolor="#" + LINE, linewidth=0.6))
        if txt == "":
            return
        tx = x + w / 2 if ha == "center" else x + pad
        ty = y + h / 2 if va == "center" else y + pad
        ax.text(tx, ty, txt, fontsize=fs, va=va, ha=ha, color=color,
                weight="bold" if bold else "normal")

    y = title_h
    cell(xs[0], y, label_w, hdr_h, "周期", "#" + GREEN_HEAD, fs=10, bold=True, ha="center", va="center")
    for di, d in enumerate(days):
        cell(xs[di + 1], y, day_w, hdr_h, d, "#" + GREEN_HEAD, fs=10.5, bold=True, ha="center", va="center")
    y += hdr_h

    cell(xs[0], y, label_w, hdr_h, "日期", "#" + GREEN_HEAD2, fs=10, ha="center", va="center")
    for di, d in enumerate(dates):
        cell(xs[di + 1], y, day_w, hdr_h, d, "#" + GREEN_HEAD2, fs=9.5, ha="center", va="center")
    y += hdr_h

    for si, slot in enumerate(slots):
        h = row_h[si]
        cell(xs[0], y, label_w, h, slot, "#" + GREEN_HEAD, fs=10, bold=True, ha="center", va="center")
        for di in range(len(days)):
            names = grid[si][di]
            zero = not names
            fc = "#" + (WARN_BG if zero else "FFFFFF")
            ax.add_patch(Rectangle((xs[di + 1], y), day_w, h, facecolor=fc,
                                   edgecolor="#" + LINE, linewidth=0.6))
            if zero:
                cx = xs[di + 1] + day_w / 2
                ax.text(cx, y + pad, "（全员有课）", fontsize=9,
                        va="top", ha="center", color="#" + WARN_TEXT)
                ax.text(cx, y + h - 0.24, "0", fontsize=10,
                        va="center", ha="center", color="#" + WARN_TEXT, weight="bold")
                continue
            lines = wrap_names(names)
            cx = xs[di + 1] + day_w / 2
            for li, ln in enumerate(lines):
                ax.text(cx, y + pad + li * line_h, ln,
                        fontsize=9, va="top", ha="center", color="#1f2328")
            ax.text(cx, y + h - 0.24, str(counts[si][di]), fontsize=10,
                    va="center", ha="center", color="#" + GREEN_TEXT, weight="bold")
        y += h

    fig.savefig(path, dpi=200, facecolor="white", bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)


# ---------------------------------------------------------------- 可视化增强
def _cjk_plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    avail = {f.name for f in font_manager.fontManager.ttflist}
    for cand in ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC", "SimSun"]:
        if cand in avail:
            plt.rcParams["font.sans-serif"] = [cand]
            break
    plt.rcParams["axes.unicode_minus"] = False
    return plt


def _mix(c1, c2, t):
    """两个 #RRGGBB 之间线性插值。"""
    a = tuple(int(c1[i:i + 2], 16) for i in (1, 3, 5))
    b = tuple(int(c2[i:i + 2], 16) for i in (1, 3, 5))
    return "#%02x%02x%02x" % tuple(round(a[i] + (b[i] - a[i]) * t) for i in range(3))


HEAT_LO, HEAT_HI = "#EAF6EC", "#2E8B57"
BEST_RING = "#E6A23C"


def render_heatmap(meta, slots, days, people, grid, dates, path):
    """时段 × 星期 热力图：格内数字 = 该时段空闲人数，颜色越深人越齐。"""
    plt = _cjk_plt()
    from matplotlib.patches import Rectangle

    counts = [[len(grid[si][di]) for di in range(len(days))] for si in range(len(slots))]
    total = len(people)
    mx = max([c for r in counts for c in r] or [0]) or 1

    label_w, col_w, cell_h = 1.55, 1.70, 0.90
    hdr_h, title_h, legend_h = 0.60, 0.68, 0.70
    x0 = 0.16
    W = x0 * 2 + label_w + col_w * len(days)
    H = title_h + hdr_h + cell_h * len(slots) + legend_h
    fig = plt.figure(figsize=(W, H), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.invert_yaxis()
    ax.axis("off")

    def box(x, y, w, h, fc, ec="white", lw=1.2):
        ax.add_patch(Rectangle((x, y), w, h, facecolor=fc, edgecolor=ec, linewidth=lw))

    def txt(x, y, s, fs, color="#1f2328", bold=False, ha="center"):
        ax.text(x, y, s, fontsize=fs, va="center", ha=ha, color=color,
                weight="bold" if bold else "normal")

    txt(x0, title_h * 0.5, f"{meta['title']}　周期 {meta['week_label']}",
        13, bold=True, ha="left")
    txt(W - x0, title_h * 0.5, f"共 {total} 人　格子越绿 = 该时段空闲的人越多",
        8.5, color="#6b7280", ha="right")

    y = title_h
    box(x0, y, label_w, hdr_h, "#" + GREEN_HEAD, ec="#" + GREEN_HEAD)
    txt(x0 + label_w / 2, y + hdr_h / 2, "时段", 10.5, "#" + GREEN_TEXT, True)
    for di, d in enumerate(days):
        cx = x0 + label_w + di * col_w
        box(cx, y, col_w, hdr_h, "#" + GREEN_HEAD, ec="#" + GREEN_HEAD)
        txt(cx + col_w / 2, y + hdr_h * 0.35, d, 10.5, "#" + GREEN_TEXT, True)
        if dates[di]:
            txt(cx + col_w / 2, y + hdr_h * 0.75, dates[di], 8.5, "#4b7a56")
    y += hdr_h

    ranked = sorted(((counts[si][di], di, si) for si in range(len(slots))
                     for di in range(len(days))), key=lambda t: (-t[0], t[1], t[2]))
    best = {(di, si) for n, di, si in ranked[:2] if n > 0}

    for si, slot in enumerate(slots):
        box(x0, y, label_w, cell_h, "#" + GREEN_HEAD, ec="#" + GREEN_HEAD)
        txt(x0 + label_w / 2, y + cell_h / 2, slot.replace(" ", ""), 9.5, "#" + GREEN_TEXT, True)
        for di in range(len(days)):
            n = counts[si][di]
            cx = x0 + label_w + di * col_w
            if n == 0:
                fc, tc, sub = "#" + WARN_BG, "#" + WARN_TEXT, "全员有课"
            else:
                t = n / mx
                fc = _mix(HEAT_LO, HEAT_HI, t)
                tc = "#ffffff" if t > 0.62 else "#1B4A28"
                sub = f"/ {total} 人"
            box(cx, y, col_w, cell_h, fc)
            txt(cx + col_w / 2, y + cell_h * 0.40, str(n), 18, tc, True)
            txt(cx + col_w / 2, y + cell_h * 0.75, sub, 7.5, tc)
            if (di, si) in best:
                ax.add_patch(Rectangle((cx + 0.025, y + 0.025), col_w - 0.05, cell_h - 0.05,
                                       fill=False, edgecolor=BEST_RING, linewidth=2.0))
        y += cell_h

    ly = y + 0.24
    txt(x0, ly, "图例", 8.5, "#6b7280", ha="left")
    lx = x0 + 0.52
    steps = 24
    for i in range(steps):
        ax.add_patch(Rectangle((lx + i * 0.09, ly - 0.085), 0.09, 0.17,
                               facecolor=_mix(HEAT_LO, HEAT_HI, i / (steps - 1)), edgecolor="none"))
    txt(lx - 0.06, ly, "0", 8, "#6b7280", ha="right")
    txt(lx + steps * 0.09 + 0.06, ly, f"{mx} 人", 8, "#6b7280", ha="left")
    gx = lx + steps * 0.09 + 0.70
    ax.add_patch(Rectangle((gx, ly - 0.085), 0.16, 0.17, facecolor="#" + WARN_BG, edgecolor="none"))
    txt(gx + 0.24, ly, "全员有课", 8, "#" + WARN_TEXT, ha="left")
    txt(W - x0, ly, "橙框 = 本周最推荐的 2 个时段", 8, "#B06F00", ha="right")

    fig.savefig(path, dpi=200, facecolor="white", bbox_inches="tight", pad_inches=0.14)
    plt.close(fig)


def render_matrix(meta, slots, days, people, grid, dates, path):
    """人 × 时段 占用矩阵：一行一人，绿=空闲 灰=有课，一眼看出谁最忙。"""
    plt = _cjk_plt()
    from matplotlib.patches import Rectangle

    n_s, n_d = len(slots), len(days)
    ncol = n_s * n_d
    name_w, col_w, row_h = 1.20, 0.40, 0.44
    hdr1, hdr2, title_h, foot_h = 0.36, 0.30, 0.66, 0.50
    cnt_w = 1.00
    x0 = 0.16
    W = x0 * 2 + name_w + col_w * ncol + cnt_w
    H = title_h + hdr1 + hdr2 + row_h * len(people) + foot_h
    fig = plt.figure(figsize=(W, H), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.invert_yaxis()
    ax.axis("off")

    def txt(x, y, s, fs, color="#1f2328", bold=False, ha="center"):
        ax.text(x, y, s, fontsize=fs, va="center", ha=ha, color=color,
                weight="bold" if bold else "normal")

    txt(x0, title_h * 0.52, f"{meta['title']}·人员占用矩阵　周期 {meta['week_label']}",
        12, bold=True, ha="left")
    txt(W - x0, title_h * 0.52, f"共 {total_n(people)} 人 × {ncol} 个时段",
        8.5, color="#6b7280", ha="right")

    y = title_h
    gx = x0
    ax.add_patch(Rectangle((gx, y), name_w, hdr1 + hdr2, facecolor="#" + GREEN_HEAD, edgecolor="white"))
    txt(gx + name_w / 2, y + (hdr1 + hdr2) / 2, "姓名", 9.5, "#" + GREEN_TEXT, True)
    for di, d in enumerate(days):
        cx = x0 + name_w + di * n_s * col_w
        ax.add_patch(Rectangle((cx, y), n_s * col_w, hdr1, facecolor="#" + GREEN_HEAD,
                               edgecolor="white", linewidth=1.0))
        label = d if not dates[di] else f"{d} {dates[di]}"
        txt(cx + n_s * col_w / 2, y + hdr1 / 2, label, 8.5, "#" + GREEN_TEXT, True)
        for si, slot in enumerate(slots):
            sx = cx + si * col_w
            ax.add_patch(Rectangle((sx, y + hdr1), col_w, hdr2,
                                   facecolor="#" + GREEN_HEAD2, edgecolor="white", linewidth=0.6))
            txt(sx + col_w / 2, y + hdr1 + hdr2 / 2, slot.replace("上午", "").replace("下午", "")
                .replace("节", "").strip(), 6.5, "#48505a")
    ax.add_patch(Rectangle((x0 + name_w + ncol * col_w, y), cnt_w, hdr1 + hdr2,
                           facecolor="#" + GREEN_HEAD, edgecolor="white"))
    txt(x0 + name_w + ncol * col_w + cnt_w / 2, y + (hdr1 + hdr2) / 2, "空闲/共", 8, "#" + GREEN_TEXT, True)
    y += hdr1 + hdr2

    for p in people:
        nfree = ncol - sum(1 for (d, s) in p["busy"] if d <= n_d and s <= n_s)
        ax.add_patch(Rectangle((x0, y), name_w, row_h, facecolor="#FAFCFA", edgecolor="white"))
        txt(x0 + name_w / 2, y + row_h / 2, p["name"], 8.5, ha="center")
        for di in range(n_d):
            for si in range(n_s):
                cx = x0 + name_w + (di * n_s + si) * col_w
                busy = (di + 1, si + 1) in p["busy"]
                ax.add_patch(Rectangle((cx + 0.015, y + 0.02), col_w - 0.03, row_h - 0.04,
                                       facecolor="#C8CFD7" if busy else "#CDEBD4",
                                       edgecolor="none"))
        ax.add_patch(Rectangle((x0 + name_w + ncol * col_w, y), cnt_w, row_h,
                               facecolor="#FAFCFA", edgecolor="white"))
        txt(x0 + name_w + ncol * col_w + cnt_w / 2, y + row_h / 2, f"{nfree}/{ncol}", 8.5,
            "#" + GREEN_TEXT if nfree >= ncol / 2 else "#8a929c", True)
        y += row_h

    ly = y + 0.24
    ax.add_patch(Rectangle((x0, ly - 0.085), 0.20, 0.17, facecolor="#CDEBD4", edgecolor="none"))
    txt(x0 + 0.26, ly, "空闲", 8.5, "#6b7280", ha="left")
    ax.add_patch(Rectangle((x0 + 0.90, ly - 0.085), 0.20, 0.17, facecolor="#C8CFD7", edgecolor="none"))
    txt(x0 + 1.16, ly, "有课", 8.5, "#6b7280", ha="left")
    txt(W - x0, ly, "右侧数字 = 本周空闲时段数，越大越闲", 8, "#6b7280", ha="right")

    fig.savefig(path, dpi=200, facecolor="white", bbox_inches="tight", pad_inches=0.14)
    plt.close(fig)


def total_n(people):
    return len(people)


# ---------------------------------------------------------------- main
TEMPLATE = {
    "meta": {"title": "实验室人员空闲表", "week_label": "8.31-9.6", "monday_date": "8.31"},
    "slots": DEFAULT_SLOTS,
    "days": DEFAULT_DAYS,
    "people": [{"name": "示例-姓名", "busy": [[1, 2], [1, 3], [3, 1]]}],
}


def main():
    ap = argparse.ArgumentParser(description="课表数据 -> 人员空闲表")
    ap.add_argument("data", nargs="?", help="数据 JSON 路径")
    ap.add_argument("-o", "--out", default=None, help="输出目录（默认与 JSON 同目录）")
    ap.add_argument("--week", default=None, help="周期标签，如 8.31-9.6")
    ap.add_argument("--monday", default=None, help="周一日期，如 8.31")
    ap.add_argument("--formats", default="html,xlsx,csv,png,heatmap,matrix",
                    help="要输出的格式，逗号分隔：html,xlsx,csv,png,heatmap,matrix")
    ap.add_argument("--name", default="人员空闲表", help="输出文件名（不含扩展名）")
    ap.add_argument("--days", type=int, default=0, help="只输出前 N 天（0=全部，常用 5 = 周一到周五）")
    ap.add_argument("--drop-empty-slots", action="store_true",
                    help="删掉没人有课的大节行（例如没人上晚课时去掉「晚上」）")
    ap.add_argument("--template", action="store_true", help="打印数据模板后退出")
    args = ap.parse_args()

    if args.template:
        print(json.dumps(TEMPLATE, ensure_ascii=False, indent=2))
        return 0

    if not args.data:
        ap.error("缺少 data.json 路径（或用 --template 打印模板）")

    meta, slots, days, people = load_data(args.data, args)

    if args.days and args.days < len(days):
        nd = args.days
        days = days[:nd]
        people = [{"name": p["name"], "busy": {(d, s) for (d, s) in p["busy"] if d <= nd}}
                  for p in people]

    if args.drop_empty_slots:
        used = {s for p in people for (_d, s) in p["busy"]}
        keep = [i for i in range(len(slots)) if (i + 1) in used]
        if keep:
            remap = {old + 1: new + 1 for new, old in enumerate(keep)}
            people = [{"name": p["name"],
                       "busy": {(d, remap[s]) for (d, s) in p["busy"] if s in remap}}
                      for p in people]
            slots = [slots[i] for i in keep]

    dates = week_dates(meta["monday_date"], days)
    grid = free_map(people, days, slots)
    outdir = args.out or os.path.dirname(os.path.abspath(args.data))
    os.makedirs(outdir, exist_ok=True)

    fmts = [f.strip().lower() for f in args.formats.split(",") if f.strip()]
    made = []
    base = os.path.join(outdir, args.name)
    if "html" in fmts:
        render_html(meta, slots, days, people, grid, dates, base + ".html")
        made.append(base + ".html")
    if "csv" in fmts:
        render_csv(meta, slots, days, people, grid, dates, base + ".csv")
        made.append(base + ".csv")
    if "xlsx" in fmts:
        render_xlsx(meta, slots, days, people, grid, dates, base + ".xlsx")
        made.append(base + ".xlsx")
    if "png" in fmts:
        try:
            render_png(meta, slots, days, people, grid, dates, base + ".png")
            made.append(base + ".png")
        except Exception as e:
            print(f"[warn] PNG 渲染失败：{e.__class__.__name__}: {e}")
    for key, fn, suffix in (("heatmap", render_heatmap, "_热力图.png"),
                            ("matrix", render_matrix, "_人员矩阵.png")):
        if key in fmts:
            try:
                fn(meta, slots, days, people, grid, dates, base + suffix)
                made.append(base + suffix)
            except Exception as e:
                print(f"[warn] {key} 渲染失败：{e.__class__.__name__}: {e}")

    print(f"人数：{len(people)}　时段：{len(slots)} 个大节 × {len(days)} 天")
    print("空闲人数矩阵（行=大节，列=星期）：")
    for si, slot in enumerate(slots):
        print("  " + slot.ljust(12) + " " + "  ".join(f"{len(grid[si][di]):>2}" for di in range(len(days))))
    active = {di for p in people for (di, _si) in p["busy"] if di <= len(days)}
    pool = active if active else set(range(1, len(days) + 1))
    best = sorted(((len(grid[si][di - 1]), di, si) for si in range(len(slots)) for di in pool),
                  key=lambda t: (-t[0], t[1], t[2]))[:3]
    print("推荐时段 TOP3（只统计有人上课的日子）："
          + "；".join(f"{days[di-1]} {slots[si]}（{n} 人）" for n, di, si in best))
    warn = check_anomalies(people, days, slots, grid)
    if warn:
        print("需核对：")
        for w in warn:
            print("  - " + w)
    print("已生成：")
    for m in made:
        print("  " + m)
    return 0


if __name__ == "__main__":
    sys.exit(main())
